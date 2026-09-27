#!/usr/bin/env python3
"""Result 3 한 장 — 편집 가능한 네이티브 차트 + 정의 표 두 개.
   입력: assets/sweep/result3_prefill_reclaim.json (scripts/result3_prefill_reclaim.py)
   출력: 260924_result3_prefill_reclaim.pptx"""
import json, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import (XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_MARK,
                             XL_LABEL_POSITION, XL_TICK_LABEL_POSITION)
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

C = lambda h: RGBColor.from_string(h)
INK, INK2, MUTED, GRID, AXIS = C('0B0B0B'), C('52514E'), C('898781'), C('E8E7E1'), C('C3C2B7')
ORANGE, TEAL, GREY, HEAD = C('EB6834'), C('1B9E77'), C('898781'), C('F3F2EE')

R = json.load(open(os.path.join(os.path.dirname(__file__), '..', 'assets', 'sweep',
                                'result3_prefill_reclaim.json')))
POL = [('admission', 'Admission only', ORANGE),
       ('reclaim',   'Admission + reclaim', TEAL),
       ('resident',  'Resident (oracle)', GREY)]

prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])

def text(x, y, w, h, runs, align=PP_ALIGN.LEFT):
    tf = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    for k, (t, sz, col, b) in enumerate(runs):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph(); p.alignment = align
        r = p.add_run(); r.text = t; r.font.size, r.font.color.rgb, r.font.bold = Pt(sz), col, b

text(.62, .34, 12.1, .9, [
    ('Result 3: Admission alone is not enough, lifetime and reclaim matter', 24, INK, True),
    ('동일한 workload 와 fast-pool 용량 (C2 2층 600 mm², pool 3.30 GB) 에서 prefill 요청 순서와 '
     'reclaim 만 바꾸고, 이어지는 decode 한 스텝의 HBM 트래픽을 본다', 11.5, MUTED, False)])

# ---------------------------------------------------------------- chart
d = CategoryChartData()
d.categories = [o['label'] for o in R['orders']]
for key, name, _ in POL:
    d.add_series(name, [round(o[key]['pct'], 2) for o in R['orders']])
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.45), Inches(1.35),
                        Inches(7.45), Inches(4.75), d).chart
ch.has_title = False; ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False
ch.legend.font.size = Pt(12); ch.legend.font.color.rgb = INK2
for ax in (ch.category_axis, ch.value_axis):
    ax.tick_labels.font.size = Pt(12); ax.tick_labels.font.color.rgb = INK2
    ax.major_tick_mark = ax.minor_tick_mark = XL_TICK_MARK.NONE
    ax.format.line.color.rgb = AXIS
ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW   # 이름은 바닥, 막대와 안 겹침
va = ch.value_axis
va.minimum_scale, va.maximum_scale, va.major_unit = -5, 25, 5
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = GRID; va.major_gridlines.format.line.width = Pt(.75)
va.tick_labels.number_format, va.tick_labels.number_format_is_linked = '0', False
va.has_title = True; va.axis_title.text_frame.text = 'HBM 물리 트래픽 감소  [%]'
for r in va.axis_title.text_frame.paragraphs[0].runs: r.font.size, r.font.color.rgb = Pt(12), INK2
ch.plots[0].gap_width, ch.plots[0].overlap = 70, 0

for ser, (_, _, col) in zip(ch.series, POL):
    ser.invert_if_negative = False                 # 없으면 PowerPoint 가 음수 막대를 흰색으로 뒤집는다
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = col
    ser.format.line.fill.background()

adm = ch.series[0]                                 # 순서에 따라 움직이는 건 이 계열뿐 — 여기만 값 표시
adm.data_labels.show_value = True
dl = adm.data_labels
dl.number_format, dl.number_format_is_linked = '0.00', False
dl.position = XL_LABEL_POSITION.OUTSIDE_END
dl.font.size = Pt(12); dl.font.bold = True; dl.font.color.rgb = ORANGE
floor = adm.points[2].data_label                  # 바닥: 물리 = 논리
floor.has_text_frame = True
floor.text_frame.text = f"{R['orders'][2]['admission']['pct']:.2f}  (물리 = 논리)"
for r in floor.text_frame.paragraphs[0].runs:
    r.font.size, r.font.bold, r.font.color.rgb = Pt(12), True, ORANGE
floor.position = XL_LABEL_POSITION.OUTSIDE_END

text(.6, 6.15, 7.3, .9, [
    ('감소 = 1 − V_HBM(정책) / V_HBM(기준)      V_HBM = V^R_HBM + V^W_HBM,  decode 정상상태 한 스텝',
     10.5, INK2, False),
    (f"기준 = 3D 없음 · 기존 L2 만 · prefill 없음 ({R['base_GB']:.2f} GB).   논리 수요는 모든 막대에서 "
     f"{R['demand_GB']:.2f} GB — 물리가 논리와 같아지면 {R['floor_pct']:.2f} %, 이 축의 바닥", 10.5, MUTED, False)])

# ---------------------------------------------------------------- tables
def table(rows, x, y, colw, rowh=.42, fs=11):
    tb = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y),
                            Inches(sum(colw)), Inches(rowh*len(rows))).table
    tb.first_row = True
    for j, w in enumerate(colw): tb.columns[j].width = Inches(w)
    for i, row in enumerate(rows):
        for j, (v, col) in enumerate(row):
            c = tb.cell(i, j); c.text = v
            c.fill.solid(); c.fill.fore_color.rgb = HEAD if i == 0 else C('FFFFFF')
            c.margin_left = c.margin_right = Inches(.06)
            p = c.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(fs); r.font.bold = (i == 0 or j == 0)
                r.font.color.rgb = col or (INK if i == 0 else INK2)
    return tb

T = lambda v, col=None: (v, col)
pol = table([
    [T('정책'), T('빈자리 있을 때'), T('가득 찼을 때'), T('해제 신호')],
    [T('Admission only', ORANGE),      T('넣음'), T('안 넣음 · 안 뺌'), T('무시')],
    [T('Admission + reclaim', TEAL),   T('넣음'), T('안 넣음 · 안 뺌'), T('즉시 비움')],
    [T('Resident (oracle)', GREY),     T('decode 가 다시 읽을 weight 만 미리 고정'), T(''), T('')],
], 8.1, 1.5, [1.55, 1.2, 1.15, 1.0], fs=10.5)
pol.cell(3, 1).merge(pol.cell(3, 3))

odr = table([
    [T('prefill 요청 순서'), T('')],
    [T(R['orders'][0]['label']), T(R['orders'][0]['desc'])],
    [T(R['orders'][1]['label']), T(R['orders'][1]['desc'])],
    [T(R['orders'][2]['label']), T(R['orders'][2]['desc'])],
], 8.1, 3.65, [1.85, 3.05])
odr.cell(0, 0).merge(odr.cell(0, 1))

text(8.1, 5.55, 4.9, .7, [
    ('티어는 데이터 종류를 모른다. 서빙 엔진이 이미 내보내는 free(buffer) 신호만 받는다.', 10.5, MUTED, False)])

out = '/Users/choeseoyeon/Desktop/DSIL/2026/09/260924_result3_prefill_reclaim.pptx'
prs.save(out); print('저장:', out)
