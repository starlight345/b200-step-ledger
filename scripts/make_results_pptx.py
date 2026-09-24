#!/usr/bin/env python3
"""More experiments + Results — 편집 가능한 네이티브 차트·표.
   출력: 260923_results_charts.pptx"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import (XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_MARK,
                             XL_MARKER_STYLE, XL_LABEL_POSITION)
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

C = lambda h: RGBColor.from_string(h)
INK, INK2, MUTED, GRID, AXIS = C('0B0B0B'), C('52514E'), C('898781'), C('E8E7E1'), C('C3C2B7')
BLUE, BLUE2, BLUE3, BLUE4, BLUE5 = C('17518F'), C('2A78D6'), C('4A8DDC'), C('7DB0E8'), C('A9C9F0')
ORANGE, TEAL, RED, GREY = C('EB6834'), C('1B9E77'), C('D03B3B'), C('898781')

prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]

def title(s, t, sub=None):
    tb = s.shapes.add_textbox(Inches(.62), Inches(.34), Inches(12.1), Inches(.85))
    tf = tb.text_frame; tf.word_wrap = True
    r = tf.paragraphs[0].add_run(); r.text = t
    r.font.size, r.font.bold, r.font.color.rgb = Pt(24), True, INK
    if sub:
        r2 = tf.add_paragraph().add_run(); r2.text = sub
        r2.font.size, r2.font.color.rgb = Pt(11.5), MUTED

def axt(ax, t):
    ax.has_title = True; tf = ax.axis_title.text_frame; tf.text = t
    for r in tf.paragraphs[0].runs: r.font.size, r.font.color.rgb = Pt(12), INK2

def quiet(ch, xt, yt, legend=True, pos=XL_LEGEND_POSITION.BOTTOM):
    ch.has_title = False; ch.has_legend = legend
    if legend:
        ch.legend.position = pos; ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(12); ch.legend.font.color.rgb = INK2
    for ax, t in ((ch.category_axis, xt), (ch.value_axis, yt)):
        ax.tick_labels.font.size = Pt(12); ax.tick_labels.font.color.rgb = INK2
        ax.major_tick_mark = XL_TICK_MARK.NONE; ax.minor_tick_mark = XL_TICK_MARK.NONE
        ax.format.line.color.rgb = AXIS
        if t: axt(ax, t)
    ch.value_axis.has_major_gridlines = True
    gl = ch.value_axis.major_gridlines.format.line; gl.color.rgb = GRID; gl.width = Pt(0.75)
    ch.category_axis.has_major_gridlines = False

def sline(ser, rgb, w=2.5, dash=None, marker=True, ms=7):
    ser.format.line.color.rgb = rgb; ser.format.line.width = Pt(w); ser.smooth = False
    if dash: ser.format.line.dash_style = dash
    if marker:
        ser.marker.style = XL_MARKER_STYLE.CIRCLE; ser.marker.size = ms
        ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb
        ser.marker.format.line.color.rgb = rgb
    else: ser.marker.style = XL_MARKER_STYLE.NONE

def bar(ser, rgb):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb
    ser.format.line.fill.background()

def table(s, rows, x, y, w, h, colw=None, fs=12.5, hl=None):
    tb = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
    if colw:
        for j, cw in enumerate(colw): tb.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = tb.cell(i, j); c.text = str(v)
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(fs); r.font.bold = (i == 0) or (hl and i in hl)
                r.font.color.rgb = INK if (i == 0 or (hl and i in hl)) else INK2
    return tb

def note(s, x, y, w, lines):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(1.6))
    tf = tb.text_frame; tf.word_wrap = True
    for k, (t, sz, col, b) in enumerate(lines):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = t
        r.font.size, r.font.color.rgb, r.font.bold = Pt(sz), col, b

# ============================================ 1. More experiments are needed
s = prs.slides.add_slide(BLANK)
title(s, 'More experiments are needed',
      '트래픽과 소자 요구는 측정으로 받쳐진다. 시간(성능)만 미검증 가정 위에 있다.')
rows = [['','무엇을 주장하나','현재 근거','막혀 있는 것'],
        ['트래픽 감소  18.6 ~ 24.8 %','HBM 물리 바이트가 준다','측정  trace replay','—'],
        ['충전  13.86 → 0 GB/step','수요 충전은 쓰고 안 읽는다','측정  trace replay','—'],
        ['요구 쓰기 BW  43.18 → 없음','소자 스펙에서 쓰기가 빠진다','측정 유도','—'],
        ['순서 불변성  ±0.01 %p','결과가 trace 성질이 아니다','측정  6개 순서','—'],
        ['speedup  1.077 ~ 1.123','시간이 그만큼 준다','projection','ΔT = ΔD / 6.40  미검증'],
        ['L2 적중률  ~ 0','decode 는 L2 를 못 쓴다','가정','카운터 미수집'],
        ['배치 64 / 128','대배치에서도 유지','외삽','B ≤ 32 만 측정'],
        ['커널 발행 순서','층 인터리빙','가정','타임라인 미수집'],
        ['70B 급 모델','평가 바닥선','없음','미실행']]
table(s, rows, .7, 1.55, 11.9, 3.9, colw=[3.5, 3.3, 2.5, 2.6], fs=12, hl=(5,))
note(s, .7, 5.65, 11.9, [
 ('이 표의 아래 다섯 줄이 한 번의 B200 세션으로 동시에 닫힌다.', 15, INK, True),
 ('실제 decode 계산은 67초다 (5조건 × 3반복 × 1,000스텝). 나머지는 전부 셋업이라 더 싼 GPU 에서 끝낸다.', 12.5, INK2, False)])

# ============================================ 2. Why B200
s = prs.slides.add_slide(BLANK)
title(s, 'Why this one experiment, and why B200',
      'HBM 바이트를 줄이면 스텝 시간이 그만큼 주는가 — 인과 계수를 직접 잰다')
rows = [['','내용'],
        ['질문','ΔT = ΔD / B_eff 의 B_eff 가 정말 6.40 TB/s 인가'],
        ['왜 아직 모르나','6.40 은 (배치, 문맥) 앵커 8점 사이의 회귀 기울기다.\n조건을 바꾸면 트래픽만이 아니라 연산량·스케줄링도 같이 바뀐다.'],
        ['방법','워크로드 완전 고정, cudaLimitPersistingL2CacheSize 만 0 ~ 상한으로 변경\ndram__bytes_read 와 step 시간을 동시 수집, 5조건 × 1,000스텝 × 3반복'],
        ['왜 B200 인가','앵커 8점이 전부 B200 · L2 126 MB · persisting 상한 82,903,040 B\n다른 기계는 그 기계의 계수가 나온다'],
        ['전제 조건','ncu 로 GPU 성능 카운터 접근 가능해야 한다 (이전 실험이 여기서 실패)'],
        ['비용','약 $28  ≈ 4 만 원 (RTX PRO 6000 에서 셋업 → B200 1시간 job)'],
        ['판정','기울기가 ±20 % 안 & R² ≥ 0.9 → 논문 수치로 승격\nR² < 0.9 → 선형 교체 모델 자체를 폐기']]
table(s, rows, .7, 1.55, 11.9, 4.3, colw=[2.4, 9.5], fs=12.5)
note(s, .7, 6.1, 11.9, [
 ('이것이 닫히지 않으면 "몇 배 빨라진다"를 한 문장도 쓸 수 없다.', 15, RED, True)])

# ============================================ 2b. 지표 정의
s = prs.slides.add_slide(BLANK)
title(s, 'Metrics', 'replay 가 직접 내는 세 값과, 거기서 정의되는 세 지표')
rows = [['','정의','단위','등급'],
        ['replay 직접 출력',
         'V^R_HBM (HBM 읽기) · V^R_tier (티어 읽기) · V^W_fill (티어 쓰기 = 충전)','GB / decode step','측정'],
        ['HBM reduction','( V^R_HBM(base) − V^R_HBM ) / V^R_HBM(base)','%','측정'],
        ['Fill cost','V^W_fill  —  정규화하지 않고 절대량 그대로','GB / decode step','측정'],
        ['Speedup','T_base / T_config','배','projection'],
        ['T_config',
         'T = t0 + [직렬: Σ 서비스 시간]  또는  [겹침: max(...)]\n'
         '서비스 시간 = V^R_HBM/B_HBM + V^R_tier/B_R + V^W_fill/B_W','ms','projection']]
table(s, rows, .7, 1.6, 11.9, 3.2, colw=[2.3, 6.1, 2.0, 1.5], fs=12.5)
note(s, .7, 5.1, 11.9, [
 ('기준선 base = 기존 B200 L2 (126 MB) 만 둔 구성. 티어를 더한 구성과 같은 logical access 를 재생한다.',
  13, INK2, False),
 ('', 8, MUTED, False),
 ('Speedup 만 projection 이다 — t0 = 1.85 ms 와 B_HBM = 6.40 TB/s 는 실측 앵커 회귀값이고, '
  '같은 워크로드에서 HBM 바이트만 줄였을 때의 인과 계수는 아직 검증되지 않았다.', 12, RED, True)])

# ============================================ 3. R1  실제 GB/step, 교차점
s = prs.slides.add_slide(BLANK)
title(s, 'Result 1: p > 0.23 부터는 아끼는 바이트보다 쓰는 바이트가 많아진다',
      '600 mm², C2 2층 · 32스텝 정상 상태, 시드 3개 · 둘 다 GB / decode step, 같은 축')
CATS = ['0','0.02','0.05','0.10','0.20','0.35','0.50','0.75','1.00']
d = CategoryChartData(); d.categories = CATS
d.add_series('이득  HBM 절감  ΔV^R', (-0.131, 3.165, 3.165, 3.165, 3.165, 3.165, 3.165, 3.165, 3.165))
d.add_series('비용  티어 충전  V^W_fill', (0.000, 0.287, 0.686, 1.392, 2.747, 4.784, 6.909, 10.359, 13.862))
ch = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(.7), Inches(1.6), Inches(8.1), Inches(4.5), d).chart
quiet(ch, 'admission 확률  p   (미스 시 티어에 할당할 확률)', '바이트  [GB / decode step]')
sline(ch.series[0], BLUE2, 3.2, ms=8); sline(ch.series[1], ORANGE, 3.2, ms=8)
ch.value_axis.minimum_scale, ch.value_axis.maximum_scale, ch.value_axis.major_unit = -2, 15, 3

rows = [['','정의','p = 1 에서'],
        ['이득  ΔV^R','V^R_HBM(base) − V^R_HBM','3.165 GB'],
        ['비용  V^W_fill','티어에 써 넣은 바이트','13.862 GB']]
table(s, rows, 8.95, 1.7, 3.75, .95, colw=[1.15, 1.6, 1.0], fs=11)
note(s, 8.95, 2.9, 3.75, [
 ('base = 기존 B200 L2 (126 MB) 만 둔 구성', 10.5, MUTED, False),
 ('V^R_HBM(base) = 17.027 GB / step', 10.5, MUTED, False),
 ('', 7, MUTED, False),
 ('이득은 p > 0 에서 3.165 GB 로 고정', 12.5, BLUE2, True),
 ('비용은 p 에 선형,  V^W_fill = 13.862 · p', 12.5, ORANGE, True),
 ('', 7, MUTED, False),
 ('교차점', 13, INK, True),
 ('p* = ΔV^R / 13.862 = 0.228', 13, INK, True),
 ('', 7, MUTED, False),
 ('p = 1 에서 비용이 이득의 4.38 배', 12.5, RED, True)])
note(s, .7, 6.25, 8.1, [
 ('충전은 미스마다 1:1 로 생긴다. 이득은 상주 집합이 차는 순간 포화하므로 더 채워도 늘지 않는다.',
  13.5, INK, True)])

# ============================================ 4. R2 최적점은 β 무관
s = prs.slides.add_slide(BLANK)
title(s, 'Result 2: 최적점은 소자 파라미터 β 와 무관하게 한 점이다',
      'β = B_W / B_R · 직렬 경계, B_R = 19 TB/s · 전량 고정(q=1)은 충전이 0 이라 β 가 식에서 빠진다')
d = CategoryChartData(); d.categories = ['0','0.02','0.05','0.10','0.20','0.35','0.50','0.75','1.00']
d.add_series('전량 고정 q=1  (β 무관)', (1.077,)*9)
d.add_series('β = 1.00', (0.995,1.073,1.067,1.058,1.041,1.016,0.991,0.953,0.917))
d.add_series('β = 0.50', (0.995,1.069,1.058,1.040,1.007,0.961,0.917,0.854,0.799))
d.add_series('β = 0.28', (0.995,1.063,1.044,1.013,0.958,0.886,0.822,0.735,0.664))
d.add_series('β = 0.10', (0.995,1.039,0.991,0.916,0.800,0.672,0.576,0.468,0.393))
d.add_series('β = 0.05', (0.995,1.004,0.918,0.798,0.637,0.489,0.394,0.299,0.240))
ch = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(.8), Inches(1.6), Inches(11.7), Inches(5.2), d).chart
quiet(ch, 'admission 확률  p', 'speedup  (티어 없음 대비)')
sline(ch.series[0], INK, 3.0, dash=4, marker=False)
for ser, col in zip(list(ch.series)[1:], (BLUE, BLUE2, BLUE3, BLUE4, BLUE5)): sline(ser, col, 2.3, ms=6)
ch.value_axis.minimum_scale, ch.value_axis.maximum_scale, ch.value_axis.major_unit = 0.2, 1.2, 0.2

# ============================================ 5. R3 순서 함정과 해법
s = prs.slides.add_slide(BLANK)
title(s, 'Result 3: 정책 하나로는 부족하다 — prefill 을 넣으면 순서에 휘둘린다',
      'prefill 프롤로그 21.59 GB 중 4.29 GB 가 한 번 읽고 버려지는 활성값 · 600 mm², C2 2층')
d = CategoryChartData()
d.categories = ['층 인터리빙\n(기본)','weight 먼저','활성 먼저\n(최악)']
d.add_series('미스에 비할당 만', (14.65, 18.59, -0.77))
d.add_series('+ 즉시 반환', (18.59, 18.59, 18.59))
d.add_series('관리 상주 (상한)', (18.59, 18.59, 18.59))
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1.1), Inches(1.6), Inches(11.1), Inches(4.4), d).chart
quiet(ch, '', 'HBM 물리 트래픽 감소  [%]', pos=XL_LEGEND_POSITION.TOP)
for ser, col in zip(ch.series, (ORANGE, TEAL, GREY)): bar(ser, col)
ch.plots[0].gap_width = 80
for ser in ch.series:
    ser.has_data_labels = True
    dl = ser.data_labels; dl.font.size = Pt(11); dl.font.color.rgb = INK2
    dl.number_format, dl.number_format_is_linked = '0.00', False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
ch.value_axis.minimum_scale, ch.value_axis.maximum_scale = -4, 22
note(s, 1.1, 6.0, 11.1, [
 ('비할당 하나만 쓰면 0 ~ 19.4 %p 로 흔들린다. 즉시 반환을 붙이면 전 순서에서 관리 상주와 같아진다.', 15, INK, True),
 ('티어는 데이터 종류를 모른다. free(buffer) 신호만 받는다 — 서빙 엔진이 이미 내보내는 신호다.', 12.5, INK2, False),
 ('Validation — decode 만 있는 장부에서는 6개 발행 순서(preserved / forward / weight-first / KV-first / '
  'operation-aware / shuffled)를 바꿔도 변화가 0.01 %p 미만이다. 모든 타일의 스텝당 상주 가치가 같기 때문이며, '
  '순서가 문제가 되는 것은 위처럼 가치가 다른 객체가 섞일 때뿐이다.', 11, MUTED, False)])

# ============================================ 6. R4 수미상관
s = prs.slides.add_slide(BLANK)
title(s, 'Result 4: 처음 질문으로 — 만들 수 있는 용량 중 실제로 몇 GB 가 트래픽을 없애나',
      '정책을 정하고 나서야 용량이 useful 해진다 · 전환율 = 실제 HBM 감소 바이트 / 티어 용량')
d = CategoryChartData(); d.categories = ['400','500','600','700','800']
d.add_series('패브릭 상한  L2 → SM  19', (19.0,)*5)
d.add_series('소자팀 덱  like-for-like  17.7', (17.7,)*5)
d.add_series('C2 2층', (7.27, 7.52, 7.79, 8.08, 8.40))
d.add_series('C3 2층', (6.95, 7.10, 7.26, 7.42, 7.59))
d.add_series('C2 1층', (6.81, 6.91, 7.03, 7.14, 7.27))
d.add_series('C3 1층', (6.66, 6.73, 6.80, 6.87, 6.95))
ch = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(.8), Inches(1.6), Inches(7.6), Inches(4.5), d).chart
quiet(ch, '순배치 면적  [mm² / die / layer]', '필요한 읽기 전달 대역폭  [TB/s]')
sline(ch.series[0], INK, 2.0, dash=4, marker=False)
sline(ch.series[1], GREY, 2.0, dash=4, marker=False)
for ser, col in zip(list(ch.series)[2:], (BLUE, BLUE3, BLUE4, BLUE5)): sline(ser, col, 2.6, ms=7)
ch.value_axis.minimum_scale, ch.value_axis.maximum_scale, ch.value_axis.major_unit = 5, 21, 2

rows = [['정책','전환율','용량 4.22 GB 가'],
        ['관리 상주 · 비할당+반환','96.9 %','4.09 GB 만큼 HBM 을 없앤다'],
        ['일반 캐시 (LRU)','−4.0 %','하나도 없애지 못한다']]
table(s, rows, 8.65, 1.75, 4.0, .95, colw=[1.9, 0.9, 1.2], fs=11)
note(s, 8.65, 3.0, 4.0, [
 ('B_min = B_HBM / (1 − ΔV^R / D)', 13, INK, True),
 ('', 6, MUTED, False),
 ('처음 숙제에서는 이 자리에 명목 용량 C 를 넣었다. 이제 replay 로 얻은 실제 감소 ΔV^R 을 넣는다.', 11.5, INK2, False),
 ('', 6, MUTED, False),
 ('모든 설계점이 덱 17.7 과 패브릭 19 아래에 있다. 최대 요구는 C2 2층 800 mm² 에서 8.40 TB/s 다.', 12, INK, True),
 ('', 6, MUTED, False),
 ('읽기 대역폭은 구속 조건이 아니다. 구속 조건은 정책이다.', 13.5, ORANGE, True)])

# ============================================ 7. Conclusion
s = prs.slides.add_slide(BLANK)
title(s, 'Conclusion', 'Turning added capacity into useful traffic reduction')
rows = [['','',''],
        ['1', 'Added capacity does not directly translate into HBM traffic reduction.',
         '일반 캐시는 20개 설계 셀 전부에서 −0.75 %, 분산 0'],
        ['2', 'Selective admission + lifetime-aware reclaim convert capacity into useful\nresidency while avoiding fill cost.',
         '충전 13.86 → 0 GB/step · 본전 쓰기 요구 43.18 TB/s → 없음'],
        ['3', 'Therefore capacity, bandwidth, and management policy must be co-designed;\nthe framework translates workload behavior into a technology target.',
         '전환율 96.9 % · 필요 읽기 BW ≤ 8.40 TB/s · 의미론 요구 둘']]
tb = table(s, rows, .7, 1.6, 11.9, 3.2, colw=[0.55, 7.15, 4.2], fs=13)
for j2 in range(3):
    c = tb.cell(0, j2); c.text = ['', 'Contribution', '근거'][j2]
    for r in c.text_frame.paragraphs[0].runs: r.font.size, r.font.bold, r.font.color.rgb = Pt(13), True, INK
note(s, .7, 5.1, 11.9, [
 ('소자 목표', 15, INK, True),
 ('전달 BW > 6.4 필수 / 15 권장 / ~19 상한   ·   용량 ≥ 4.2 GB   ·   E/bit ≤ 0.2 pJ   ·   BEOL 페리만', 13, INK2, False),
 ('의미론 ①  미스에 무조건 할당하지 말 것        의미론 ②  해제 신호를 받아 즉시 반환할 것', 13.5, ORANGE, True),
 ('', 8, MUTED, False),
 ('Limitation — 트래픽 결과는 replay 로 얻었고, end-to-end latency 개선은 projection 이다. '
  'B200 인과 검증(ΔT = ΔD / 6.40)이 남아 있다.', 11.5, MUTED, False)])

out = '/Users/choeseoyeon/Desktop/DSIL/2026/09/260923_results_charts.pptx'
prs.save(out); print('저장:', out)
