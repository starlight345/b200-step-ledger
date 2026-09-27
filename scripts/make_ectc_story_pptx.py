#!/usr/bin/env python3
"""ECTC 초록 미팅 덱 — 3D SRAM 열/전기 특성 분석. 초록 7문단을 한 장씩 show & prove.

랩 ECTC 초록 미팅 덱(정박사님 버전)을 템플릿으로 쓴다: 표지·Agenda·제목/Novelty 슬라이드는
그 자리에서 글만 바꾸고, 나머지는 같은 레이아웃(2_사용자 지정 레이아웃) 위에 새로 만든다.
근거 슬라이드 형식도 레퍼런스의 '선행 연구결과' 장을 따른다 — 굵은 소제목, 그림 하나,
하단 남색 띠에 결론 한 줄. 그 위에 증명하려는 초록 문장을 인용한다.

    python3 scripts/make_ectc_story_pptx.py <template.pptx> <out.pptx>
그림은 assets/figures/bw-*.png (scripts/make_bw_figures.py).
"""
import copy, os, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.opc.packuri import PackURI
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, '..', 'assets', 'figures')
DATE = '2026-09-24'

C = lambda h: RGBColor.from_string(h)
NAVY, INK, INK2, MUTED, LINE, WASH = C('004187'), C('000000'), C('333333'), C('7F7F7F'), C('404040'), C('E6E6E6')

src, out = sys.argv[1], sys.argv[2]
prs = Presentation(src)
SW = prs.slide_width

# ---- 구조 먼저: 표지(1)·Agenda(2)·제목/Novelty(5)만 남긴다 -------------------------------
old = list(prs.slides)
keep = {0, 1, 4}
sldIdLst = prs.slides._sldIdLst
for i, sid in enumerate(list(sldIdLst)):
    if i not in keep:
        prs.part.drop_rel(sid.rId); sldIdLst.remove(sid)
s_title, s_agenda, s_novel = old[0], old[1], old[4]
for k, sl in enumerate((s_title, s_agenda, s_novel), 1):   # 새 슬라이드와 파트 이름이 겹치지 않게
    sl.part.partname = PackURI(f'/ppt/slides/slide{k}.xml')
LAYOUT = s_novel.slide_layout
DATE_SP = next(sh for sh in s_novel.shapes if sh.is_placeholder and sh.placeholder_format.idx == 10)._element


def set_runs(par, text, keep_run=0):
    """문단의 서식을 첫 run 에서 물려받고 글만 바꾼다."""
    runs = par.runs
    runs[keep_run].text = text
    for k, r in enumerate(runs):
        if k != keep_run: r._r.getparent().remove(r._r)
    return runs[keep_run]


def set_date(slide):
    for sh in slide.shapes:
        if sh.is_placeholder and sh.placeholder_format.idx == 10:
            set_runs(sh.text_frame.paragraphs[0], DATE)


# ---- 1. 표지 ---------------------------------------------------------------------------
for sh in s_title.shapes:
    if sh.name == '부제목 2':
        ps = sh.text_frame.paragraphs
        set_runs(ps[0], '최서연')
        set_runs(ps[3], ' ')      # 이메일 자리 — 본인 주소로 채울 것
    if sh.name == '제목 1':
        set_runs(sh.text_frame.paragraphs[0], 'IEEE ECTC 2027 초록 주제 미팅')

# ---- 2. Agenda -------------------------------------------------------------------------
set_date(s_agenda)
for sh in s_agenda.shapes:
    if sh.name == 'TextBox 4':
        ps = sh.text_frame.paragraphs
        set_runs(ps[0], '초록 제목 및 Novelty of Work')
        set_runs(ps[2], 'Show & Prove — 초록 일곱 문단과 그 근거')
        set_runs(ps[4], '남은 항목 및 일정')

# ---- 3. 제목 및 Novelty ------------------------------------------------------------------
set_date(s_novel)
for sh in s_novel.shapes:
    tf = sh.text_frame if sh.has_text_frame else None
    if sh.name == '제목 5':
        set_runs(tf.paragraphs[0], '초록주제) 제목 및 Novelty of Work')
    elif sh.name == 'Rectangle 1':
        set_runs(tf.paragraphs[0], 'Duty-Cycle-Aware Thermal and Energy Budgeting')
        set_runs(tf.paragraphs[1], 'for BEOL-Stacked SRAM on LLM Inference GPUs')
        sh.left, sh.width = Inches(0.5), SW - Inches(1.0)
    elif sh.name == 'TextBox 12':
        set_runs(tf.paragraphs[0], '1st : Thermal/Mechanical Simulation & Characterization, 2nd:  Packaging Technologies')
    elif sh.name == 'TextBox 31':
        set_runs(tf.paragraphs[0], 'Novelty of work  (42 / 50 words)')
        set_runs(tf.paragraphs[1],
                 "We derive a BEOL-stacked SRAM tier's thermal and energy specification from measured "
                 "LLM-decode traffic rather than an assumed power map, and show that the standard "
                 "steady-state screening rule and its obvious transient correction miss the verified "
                 "answer by 4× in opposite directions.")
    elif sh.name == 'TextBox 9':
        ps = tf.paragraphs
        set_runs(ps[1], '3D SRAM 티어의 타당성은 관행적으로 정상상태 식 BW ≤ P / E_bit 로 거름.')
        set_runs(ps[2], '그러나 decode 티어는 스텝의 4.80%만 일하는 펄스 부하 → 열 모델마다 판정이 4배씩 갈림.')
        set_runs(ps[3], '측정된 decode 트래픽에서 티어의 열·에너지 스펙을 유도하고, 이를 안전하게 내는 '
                        '축약 열모델(≈10노드, 티어–기판 분리)을 Ansys MAPDL로 검증해 제시함.')
        ps[4]._p.getparent().remove(ps[4]._p)


# ---- 새 슬라이드 공용 -------------------------------------------------------------------
def new_slide(title, notes=None):
    s = prs.slides.add_slide(LAYOUT)
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx == 13: ph._element.getparent().remove(ph._element)
    s.shapes.title.text = title
    s.shapes._spTree.append(copy.deepcopy(DATE_SP))
    set_date(s)
    if notes: s.notes_slide.notes_text_frame.text = notes
    return s


def text(s, x, y, w, h, runs, size=14, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=None):
    """runs: 문단 목록. 문단은 문자열이거나 (문자열, dict(bold, color, size, italic, underline)) 목록."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    if margin is not None:
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    for k, para in enumerate(runs):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        if isinstance(para, str): para = [(para, {})]
        for t, st in para:
            r = p.add_run(); r.text = t
            f = r.font; f.size = Pt(st.get('size', size)); f.bold = st.get('bold', False)
            f.italic = st.get('italic', False); f.underline = st.get('underline', False)
            f.color.rgb = st.get('color', INK)
        if 'space' in (para[0][1] if para else {}): p.space_after = Pt(para[0][1]['space'])
    return tb


def subtitle(s, t):
    text(s, 0, 0.73, 13.33, 0.42, [[(t, dict(bold=True, size=18))]])


def quote(s, para, t):
    text(s, 0.1, 1.13, 13.1, 0.4, [[(f'초록 {para}  ', dict(bold=True, size=12, color=NAVY)),
                                    (f'“{t}”', dict(italic=True, size=12, color=INK2))]])


def banner(s, t, y=6.38):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(y), SW, Inches(0.46))
    sh.fill.solid(); sh.fill.fore_color.rgb = NAVY; sh.line.fill.background(); sh.shadow.inherit = False
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = '✓  ' + t
    r.font.size, r.font.bold, r.font.color.rgb = Pt(16), True, C('FFFFFF')


def figure(s, name, x=0.4, y=1.6, w=12.5, h=4.65):
    path = os.path.join(FIG, name)
    iw, ih = Image.open(path).size
    sc = min(w / iw, h / ih); fw, fh = iw * sc, ih * sc
    return s.shapes.add_picture(path, Inches(x + (w - fw) / 2), Inches(y + (h - fh) / 2), Inches(fw), Inches(fh))


def box(s, x, y, w, h, fill=None, line=LINE, lw=1.0, shape=MSO_SHAPE.RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh


def label_in(sh, t, size=12, color=INK, bold=False):
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = t; r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold


def arrow(s, x1, y1, x2, y2, color=LINE, w=1.5):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color; ln.line.width = Pt(w)
    le = ln.line._get_or_add_ln()
    tail = le.makeelement('{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd', {'type': 'triangle'})
    le.append(tail)
    return ln


def table(s, rows, x, y, w, colw, rowh=0.42, fs=12, hl_col=None):
    tb = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(rowh * len(rows))).table
    for j, cw in enumerate(colw): tb.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tb.rows[i].height = Inches(rowh)
        for j, v in enumerate(row):
            c = tb.cell(i, j); c.text = v
            c.margin_left = c.margin_right = Inches(0.08); c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY if i == 0 else (C('DCE6F2') if hl_col == j else (C('FFFFFF') if i % 2 else C('F2F2F2')))
            for p in c.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if j == 0 or i > 0 else PP_ALIGN.CENTER
                for r in p.runs:
                    r.font.size = Pt(fs); r.font.bold = (i == 0 or j == 0)
                    r.font.color.rgb = C('FFFFFF') if i == 0 else INK
    return tb


# ---- 4. 모식도 및 주요 데이터 -------------------------------------------------------------
s = new_slide('초록주제) 모식도 및 주요 데이터',
              '왼쪽: flip-chip 스택, 방열은 위(싱크) 방향. SRAM 티어는 로직 BEOL 위에 모놀리식으로 올라가므로 '
              '로직 평면보다 열 경로의 하류에 있다. 두께는 비례가 아니다. 층 목록은 '
              'scripts/thermal_stack_solver.py build_stack(). 아래 띠: decode 스텝 4.515 ms 중 티어가 일하는 0.217 ms.')
text(s, 0.1, 0.93, 3, 0.4, [[('모식도', dict(bold=True, size=18))]])
text(s, 9.1, 0.93, 3.6, 0.4, [[('주요 데이터', dict(bold=True, size=18))]], align=PP_ALIGN.CENTER)
X0, W0 = 1.3, 4.2
layers = [  # (이름, 높이, 채움, 오른쪽 주석)
    ('Heat sink', 0.38, C('BFBFBF'), '싱크 — h 역산, 접합 100 °C 동작점'),
    ('TIM', 0.20, C('D9D9D9'), '50 µm'),
    ('Si substrate', 0.80, C('F2F2F2'), '500 µm'),
    ('Logic (FEOL)', 0.26, C('7F7F7F'), '1 µm · 349 W / die (실측 698.7 W / 2)'),
    ('BEOL', 0.45, C('D9D9D9'), '8 µm · 저-k + Cu 배선'),
    ('SRAM tier 1', 0.25, NAVY, 'IGZO 티어 50 nm'),
    ('ILD', 0.21, C('D9D9D9'), 'SiO₂ ILD 300 nm'),
    ('SRAM tier 2', 0.25, NAVY, '2층 · 4.22 GB · 20 W 예산'),
    ('ILD', 0.21, C('D9D9D9'), ''),
]
y = 1.5
for name, h, fill, note in layers:
    b = box(s, X0, y, W0, h, fill=fill, lw=0.75)
    dark = fill in (NAVY, C('7F7F7F'))
    label_in(b, name, size=11 if h > 0.15 else 9, color=C('FFFFFF') if dark else INK, bold=dark)
    if note:
        text(s, X0 + W0 + 0.12, y + h / 2 - 0.16, 3.2, 0.32, [[(note, dict(size=11 if h > 0.3 else 10, color=INK2))]],
             anchor=MSO_ANCHOR.MIDDLE, margin=0)
    y += h
for k in range(7):   # 범프
    box(s, X0 + 0.25 + k * 0.58, y + 0.04, 0.18, 0.18, fill=C('A6A6A6'), lw=0.5, shape=MSO_SHAPE.OVAL)
text(s, X0 + W0 + 0.12, y - 0.02, 3.2, 0.3, [[('범프 → 패키지 (단열 가정)', dict(size=11, color=INK2))]], margin=0)
arrow(s, X0 - 0.35, 4.6, X0 - 0.35, 1.6, color=C('C00000'), w=2.25)
text(s, 0.05, 2.7, 0.9, 0.9, [[('열', dict(size=12, bold=True, color=C('C00000')))],
                              [('→ 싱크', dict(size=11, color=C('C00000')))]], align=PP_ALIGN.CENTER)
# decode 스텝 띠
yb = 5.05
text(s, X0, yb - 0.05, 6, 0.3, [[('decode 한 스텝 (Llama-3.1-8B, B200 장부)', dict(size=11, bold=True))]], margin=0)
for k in range(3):
    x = X0 + k * 1.95
    box(s, x, yb + 0.3, 1.95, 0.34, fill=C('FFFFFF'), lw=0.75)
    box(s, x, yb + 0.3, 0.1, 0.34, fill=INK, line=None)
text(s, X0, yb + 0.68, 6.0, 0.55,
     [[('■ 티어 버스트 0.217 ms', dict(size=11, bold=True)), ('  /  스텝 4.515 ms  →  듀티 4.80%', dict(size=11))]],
     margin=0)
items = [('듀티 4.80%', '측정된 decode 트래픽 장부에서 유도\n(가정한 power map 아님)'),
         ('5.0 / 18.6 / 77.3 TB/s', '같은 소자에 세 열 모델\n정상상태 · 검증값 · 단일노드 럼프드'),
         ('MAPDL 0.02% · 1.3%', '1D 과도 솔버를 Ansys MAPDL\n1D · 3D로 교차검증'),
         ('0.019–0.260 < 0.391 pJ/bit', '문헌 기반 E/bit 예산이\n정상상태 기준선 아래 (1.5배)')]
for k, (big, small) in enumerate(items):
    text(s, 8.6, 1.55 + k * 1.2, 4.6, 1.1, [[(big, dict(size=20, bold=True, color=NAVY))],
                                             [(small, dict(size=13, color=INK2))]], align=PP_ALIGN.CENTER)

# ---- 5. Show & Prove 지도 ----------------------------------------------------------------
s = new_slide('Show & Prove — 초록 일곱 문단과 그 근거',
              '초록은 문단마다 일이 하나다(2026-09-24 재구성). 이후 슬라이드는 이 표 순서대로 한 문단의 주장을 '
              '인용하고 그것을 보이는 그림 하나를 붙인다.')
subtitle(s, '문단마다 주장 하나, 주장마다 그림 하나')
rows = [['문단', '역할', '초록이 하는 주장', '근거', '쪽'],
        ['P1', 'Problem', '관행 식 BW ≤ P/E_bit 은 이 부하를 기술하지 않는다 — 판정이 4배씩 틀림', '관행 식 판정', '6'],
        ['P2', 'Load', 'decode 티어는 한 스텝의 4.80%만 일한다 (prefill < 0.5%)', '측정 장부 · 운전점 5개', '7–8'],
        ['P3', 'Three answers', '정상상태 5.0 / 럼프드 77.3 / 층상 과도 18.6 TB/s — 10노드면 충분', '1D 과도 솔버 · 차수 스윕', '9–10'],
        ['P4', 'Verification', '해석해 2건 + MAPDL 1D 동일 18.6, 3D 균일 1.3%, 집중 시 502배는 상한', 'Ansys MAPDL 26.1', '11–13'],
        ['P5', 'Not heat', '티어는 접합 대비 0.43 K — BEOL 95% 대 소자 0.3%, 32 코너 모두 통과', '민감도 · 코너 탐색', '14–16'],
        ['P6', 'Nor energy', 'E/bit 0.019–0.260 < 0.391 pJ/bit; 남은 리스크는 읽기 마진 = 셀 설계 문제', '문헌 예산 · σVth', '17–18'],
        ['P7', 'Contribution', '측정 트래픽에서 열·에너지 스펙 유도 + 이를 안전하게 내는 축약 모델 식별', '선행 대비', '19']]
table(s, rows, 0.45, 1.35, 12.4, [0.7, 1.7, 6.9, 2.4, 0.7], rowh=0.56, fs=13)
banner(s, '스크리닝 식이 틀린 이유 → 제대로 푼 값 → 검증 → 무엇이 막는가 → 기여', y=6.2)

# ---- 근거 슬라이드 -----------------------------------------------------------------------
EVID = [
    ('① Problem — 관행 스크리닝의 판정', 'P1', 'bw-setup.png',
     '관행 스크리닝 식 BW ≤ P / E_bit 은 이 티어를 기각한다',
     '…its verdict can be wrong by a factor of four in either direction.',
     '20 W · 0.5 pJ/bit → 5.0 TB/s < HBM 실질 6.40 TB/s — 관행 식대로면 만들 이유가 없다',
     '한 스텝에서 읽는 바이트 17.16 GB(weight 15.01 + KV 2.15), 스텝 4.515 ms 중 2.66 ms가 메모리 대기. '
     '소자 모델 예시값 0.5 pJ/bit, 20 W 티어 예산. 근거: assets/sweep/canonical_constants.json.'),
    ('② Load — 듀티를 워크로드에서 유도', 'P2', 'bw-duty.png',
     '티어는 한 스텝의 4.80%만 일한다 — 정상상태 부하가 아니다',
     'the tier is active only 4.80% of the time. We call this fraction the duty cycle.',
     '4.12 GB ÷ 19 TB/s = 0.217 ms,  ÷ 4.515 ms = 듀티 4.80%',
     '설계점 C2 2층 4.22 GB가 스텝 수요의 24%를 서비스. 전달 대역폭은 온다이 패브릭 상한 19 TB/s. '
     '근거: scripts/ectc_thermal_model.py.'),
    ('② Load — 운전점을 바꿔도', 'P2', 'bw-envelope.png',
     '듀티는 한 점에서 뽑은 수가 아니다 — 운전점에 거의 무관하고 prefill은 더 낮다',
     'Decode is the worst case; in prefill the duty cycle stays below 0.5%.',
     '실측 운전점 5개에서 듀티 변동 5.1% · B200 환산 4.77% ≈ 설계점 4.80% · prefill 0.08–0.41%',
     'RTX PRO 5000 Blackwell 실측(G-5). burst와 step이 둘 다 수요 D에 비례해 상쇄되므로 듀티는 f·B_eff/B_R로 수렴. '
     'prefill은 같은 상주 바이트를 더 긴 스텝에 나르므로 듀티가 10배 이상 낮다(G-4: 로직 전력도 +4.1%뿐). '
     '근거: g5_operating_points.py, prefill_phase.py.'),
    ('③ Three answers — 같은 소자, 세 판정', 'P3', 'bw-bracket.png',
     '같은 소자 · 같은 20 W · 같은 0.5 pJ/bit 에 열 모델만 바꾸면',
     'A layered one-dimensional transient model gives 18.6 TB/s: the tier is viable…',
     '정상상태 3.7배 과소 · 단일노드 럼프드 4.16배 과대 → 검증값 18.6 TB/s ≈ 패브릭 상한 19',
     'peak fraction: 정상상태 1.0000 / 럼프드 0.0647 / 1D 솔버 0.2693. 스택 전체 τ는 6.8 ms지만 '
     '티어는 4.5 ms 주기 안에서 국소적으로 데워진다. 근거: scripts/thermal_stack_solver.py '
     '(주의: 초록 본문에는 럼프드가 76.9로 남아 있음 — 현재 솔버 값은 77.3).'),
    ('③ Three answers — 왜 럼프드가 틀리나', 'P3', 'bw-order.png',
     '럼프드가 틀리는 이유는 해상도가 아니라 위상이다',
     'about ten nodes suffice, provided tier and substrate are kept separate.',
     '축약 모델은 ~10노드면 3.8% — 조건은 티어와 기판을 한 노드로 묶지 않는 것',
     '솔버·경계조건·전력파형 고정, 메시만 변경. 1노드는 τ를 정답으로 받아도 4.16배, 기하에서 쌓으면 4.83배 낙관. '
     '물리 층당 1노드(8)면 25%, 10노드 3.8%, 16노드 0.61%. CFP의 model order reduction 항목. 근거: scripts/model_order.py.'),
    ('④ Verification — 네 경로로 검증', 'P4', 'bw-verify.png',
     '같은 스택을 해석해 2건과 Ansys MAPDL 1D · 3D로 확인했다',
     '…and Ansys MAPDL on the identical stack, which returns the same 18.6 TB/s.',
     'MAPDL 26.1 peak fraction 0.2693 대 0.2693 (0.02%) — 럼프드 오차 4.16배를 독립 재현',
     'V-1: 서버 dsil-sy의 RedHawk-SC ET 번들 MAPDL 26.1, -np 1. 계수 6.40 TB/s의 인과성은 자체 GPU에서 두 손잡이로 '
     '실측(G-1). 근거: mapdl_correlation.py, g1_causal_gpu.py.'),
    ('④ Verification — 3D 대표 모델', 'P4', 'bw-mapdl.png',
     '3D MAPDL: 349 W 로직 위에 티어를 얹은 대표 모델',
     'With the tier power spread uniformly, which is the decode case…, it agrees with the 1D model to 1.3%.',
     '싱크 93.6 → 접합 100.00 → 티어 100.01 °C: 티어 자기발열 7.5 mK · BEOL 이방성에 7자리까지 무관',
     '싱크 h = 8187 W/m²K를 접합 100 °C 동작점에 역산(캘리브레이션). 균일 전력에는 횡기울기가 없어 kxy/kz 1·10·50 모두 '
     '0.8476753 K. 근거: scripts/mapdl_representative.py, assets/mapdl/png/.'),
    ('④ Verification — 1D로 충분한가', 'P4', 'bw-hotspot.png',
     '같은 전력을 좁게 모으면 저-k BEOL이 대가를 문다',
     'Concentrating the same power into a 100 µm patch instead raises the peak 502×… That is an upper bound',
     '100 µm 패치 425 K(502배)는 등방 최악 상한 — decode는 전 매크로에서 읽으므로 균일이 동작점',
     '전력 총량 고정, 켜지는 면적만 축소. 1/면적 스케일 — 원인은 티어 아래 8 µm 저-k BEOL. 현실적 배선 이방성 10~50배면 '
     '385배·235배로 내려간다. 근거: scripts/mapdl_hotspot.py.'),
    ('⑤ What limits the tier — 열이 아니다', 'P5', 'bw-layers.png',
     '층수를 늘려도 티어 온도는 거의 그대로다',
     "the tier's peak stays within 0.43 K of the junction for any tier count.",
     '최악 BEOL k에서도 티어 peak는 접합 + 0.43 K, 층수 무관 — 층수를 멈추는 것은 열이 아니다',
     '로직 349 W/die, 접합 100 °C. 티어 평균 전력 3.65 W = 패키지 698.7 W의 0.52%. 왼쪽(이득 포화)은 입장 판정 결과로 '
     '논문 B 소관 — 여기서는 "열이 층수를 제한하지 않는다"만 쓴다. 근거: ectc_layer_sweep.py.'),
    ('⑤ What limits the tier — 민감도', 'P5', 'bw-sensitivity.png',
     '열의 주인은 소자가 아니라 BEOL 집적이다',
     'the surrounding BEOL moves the result by 95% and the tier material by 0.3%',
     'BEOL k 95% · 두께 73% 대 소자 a-IGZO k 0.3% — 완화책도 집적 선택 (둘 합쳐 1.86배)',
     '기판 박막화 500→200 µm는 1.04배뿐: R_th는 낮아지지만 듀티 평균을 해주던 열용량이 사라진다(정상상태 연구였다면 '
     '과대평가했을 항목). 근거: scripts/thermal_sensitivity.py.'),
    ('⑤ What limits the tier — 32 코너', 'P5', 'bw-corners.png',
     '측정 안 된 다섯 값을 동시에 최악으로 두어도',
     'all 32 combinations still outperform the HBM.',
     '32 / 32 코너가 HBM 6.40 TB/s를 넘는다 — 최악 7.51 TB/s (1.17배)',
     'IGZO k, BEOL k, ILD 두께, BEOL 두께, 기판 두께의 양 끝 2^5 완전탐색. 1.17배는 E/bit 0.5 예시값 기준이며 '
     '문헌 밴드를 쓰면 여유는 훨씬 크다. 근거: scripts/uncertainty_envelope.py.'),
    ('⑥ Nor energy — E/bit 예산', 'P6', 'bw-ebit.png',
     '에너지도 막지 않는다: 문헌 기반 E/bit 예산이 기준선 아래',
     'a literature-based budget of 0.019–0.260 pJ/bit stays below the 0.391 pJ/bit…',
     '0.019–0.260 < 0.391 pJ/bit — 가장 보수적인 정상상태 기준으로도 비관 코너 1.5배 여유',
     '공개 5 nm Si SRAM 어레이 에너지 × gain-cell 에너지비(6T 대비 0.5배, 비관 2배). 0.391은 20 W에서 HBM 6.40을 '
     '정상상태로 내리는 값. 근거: scripts/ebit_budget.py.'),
    ('⑥ Device risk — 읽기 마진', 'P6', 'bw-t3.png',
     '남은 소자 리스크는 읽기 마진 — 공개 σVth 20–40 mV로 다시 풀면',
     'which makes it a cell-design question rather than an open one.',
     '감당 가능한 S_max 0.27–1.08이 6T 셀 구간(0.3–1.0) 안 — 소자팀에 남은 요청은 dσ/dT 하나',
     '4.22 GB 어레이, 수율 99% → 7.20σ 필요. 판정식 S·σ·√k·nσ ≤ 155 mV(덱 read SNM). σ는 imec 300 mm IGZO(Mitard 외 2020), '
     '웨이퍼 스케일이라 국소 mismatch의 보수적 상한. 근거: scripts/t3_vth_margin.py.'),
]
for title, para, fig, sub, q, ban, note in EVID:
    s = new_slide(title, note)
    subtitle(s, sub)
    quote(s, para, q)
    figure(s, fig)
    banner(s, ban)

# ---- 19. Contribution ------------------------------------------------------------------
s = new_slide('⑦ Contribution — 선행 대비 위치',
              'imec IEDM 2025 17-3 (3D HBM-on-GPU thermal STCO): power map 가정, training, 시변 전력 언급 없음. '
              'CMOS+X (GT): 산화물 트랜지스터 메모리를 GPGPU에 적층, 전문에 열·온도 언급 0건. 기여 수준은 중간이며 강하게 쓰지 않는다.')
subtitle(s, '듀티 과도 열 해석 자체는 교과서 — 새로운 것은 부하의 출처와 모델 선택')
quote(s, 'P7', 'This work derives the tier’s thermal and energy specification from measured decode traffic, '
               'and identifies which reduced-order thermal model can safely produce it.')
rows = [['', 'imec, IEDM 2025', 'CMOS+X (Georgia Tech)', '본 연구'],
        ['메모리', '마이크로범프 HBM 스택', '산화물 TR 메모리 on GPGPU', '모놀리식 BEOL SRAM 티어'],
        ['워크로드', 'Training', 'Rodinia 벤치마크', 'LLM decode (+ prefill 확인)'],
        ['전력 입력', '가정한 power map', '—', '측정된 트래픽 장부에서 유도'],
        ['시간 구조', '언급 없음', '—', '듀티 4.80%가 주장 전체'],
        ['열 해석', '완화 (어떻게 식힐까)', '없음', '판정 (거르는 기준이 맞나) + 축약 모델'],
        ['검증', '—', '—', '해석해 2 + MAPDL 1D·3D + GPU 실측 계수']]
table(s, rows, 0.6, 1.75, 12.1, [1.6, 3.2, 3.2, 4.1], rowh=0.55, fs=14, hl_col=3)
banner(s, '측정된 decode 트래픽 → 티어의 열·에너지 스펙, 그리고 그것을 안전하게 내는 축약 열모델')

# ---- 20. 남은 항목 및 일정 ------------------------------------------------------------------
s = new_slide('남은 항목 및 일정',
              '초록 마감 2026-10-05. 본문 원고 ~2027-02. T-1/T-2는 문헌 앵커로 초록 단계는 닫혀 있음. '
              '초록 본문 수정 필요 2건: 럼프드 76.9 → 77.3 TB/s, 초록 그림(ectc-abstract)의 E/bit 밴드 0.014 → 0.019.')
subtitle(s, '초록은 문헌 앵커로 닫혀 있고, 남은 것은 측정 요청과 원고')
rows = [['#', '항목', '상태', '다음 행동'],
        ['T-1', 'E/bit 실측 (어레이 + 페리 + MIV)', '문헌 앵커, 비관 코너 1.5배 여유', '소자팀 실측값으로 gain-cell 비 대체'],
        ['T-2', '실제 스택 단면', '문헌 앵커 (ILD 300 nm, IGZO 6 nm)', '결론은 2% 내 불변 — 원고에서 교체'],
        ['T-3', 'BEOL 소자 Tj 한계 / 읽기 마진', 'σVth 20–40 mV로 셀 설계 문제로 축소', '소자팀에 dσ/dT 요청'],
        ['—', '초록 숫자 동기화', '럼프드 76.9 → 77.3, 그림 밴드 0.014 → 0.019', '제출 전 수정']]
table(s, rows, 0.6, 1.4, 12.1, [0.8, 3.6, 3.9, 3.8], rowh=0.55, fs=13)
text(s, 0.6, 4.95, 12.1, 1.2,
     [[('마감  ', dict(bold=True, size=16, color=NAVY)), ('2026-10-05  ·  700단어 + Novelty 50단어 + 그림 1장', dict(size=16))],
      [('분량  ', dict(bold=True, size=16, color=NAVY)), ('초록 700 / 700  ·  Novelty 42 / 50', dict(size=16))],
      [('세션  ', dict(bold=True, size=16, color=NAVY)), ('1st Thermal/Mechanical Simulation & Characterization  ·  2nd Packaging Technologies', dict(size=16))]])
banner(s, '열·에너지는 티어를 막지 않는다 — 남은 질문은 소자의 dσ/dT 하나')

prs.save(out)
print('wrote', out, len(prs.slides), 'slides')
